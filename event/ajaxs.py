from participant.models import Participant, StatusParticipant
from utils.utils import *
from app.decorators import *
import numpy as np

import json
import openpyxl

from django.http import HttpResponse

from openpyxl.styles import Font

from django.http import JsonResponse, HttpResponse
from django.core import serializers
from django.core.serializers.json import DjangoJSONEncoder
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger

from django.db.models import Q, BooleanField, CharField, Value, Sum, F, Case, When, Exists, Prefetch, Count
from django.db.models.functions import Coalesce, Concat
from django.db.models.fields.files import ImageFieldFile
from django.forms.models import model_to_dict
from django.forms import modelformset_factory

from django_tables2 import RequestConfig

from django.template import RequestContext
from django.template.loader import render_to_string

from datetime import datetime, date
from django.views.generic import ListView

from django.utils.safestring import mark_safe
from django.contrib.postgres.search import SearchVector
from decimal import Decimal
from django.db import IntegrityError, transaction

#EMAIL VERIFICATION AND TOKEN
from django.contrib.sites.shortcuts import get_current_site
from django.utils.http import urlsafe_base64_encode, urlsafe_base64_decode
from django.core.mail import EmailMessage

import hashlib
from uuid import uuid4
from rest_framework.authtoken.models import Token
from django.views.decorators.csrf import csrf_exempt
from authentication.models import *
from form.models import *
from authentication.forms import *

from banner.models import *
import pandas as pd
import json
from banner.filters import *
# from banner.ads_forms import *

#Set Loggin
import logging
logger = logging.getLogger(__file__)

def filter_gallery(request):
    if request.method == 'POST':
        gallery_query = Gallery.objects.all().order_by('-id')

        page = request.POST.get('page')
        per_page = request.POST.get('per_page')
        sort_by = request.POST.get('sort_by')

# @csrf_exempt
# def upload_participant(request):
#     if request.method == "POST":

#         uploaded_file = request.FILES.get("excelFile")

#         if not uploaded_file:
#             return JsonResponse({
#                 "status": "failed",
#                 "message": "File tidak ditemukan"
#             })

#         try:
#             file_name = uploaded_file.name.lower()

#             # =========================
#             # READ EXCEL / CSV
#             # =========================
#             if file_name.endswith(".csv"):
#                 df = pd.read_csv(uploaded_file)

#             elif file_name.endswith(".xlsx"):
#                 df = pd.read_excel(
#                     uploaded_file,
#                     engine="openpyxl"
#                 )

#             elif file_name.endswith(".xls"):
#                 df = pd.read_excel(
#                     uploaded_file,
#                     engine="xlrd"
#                 )

#             else:
#                 return JsonResponse({
#                     "status": "failed",
#                     "message": "Format file tidak didukung"
#                 })

#             # convert dataframe to list json
#             df = df.fillna("")

#             # convert timestamp/date columns
#             for col in df.columns:
#                 df[col] = df[col].apply(
#                     lambda x: x.strftime("%Y-%m-%d %H:%M:%S")
#                     if isinstance(x, pd.Timestamp)
#                     else x
#                 )

#             excel_json = df.to_dict(orient="records")

#             # ambil form
#             form_obj = Form.objects.first()

#             # =========================
#             # SAVE RAW JSON
#             # =========================
#             form_response = FormResponse.objects.create(
#                 form=form_obj,
#                 data_response=excel_json
#             )

#             # =========================
#             # SAVE PARTICIPANTS
#             # =========================
#             for row in excel_json:
#                 Participant.objects.create(
#                     form=form_obj,
#                     name=row.get("name") or row.get("Name"),
#                     email=row.get("email") or row.get("Email"),
#                     phone=row.get("phone") or row.get("Phone")
#                 )

#             return JsonResponse({
#                 "status": "success",
#                 "message": "Upload berhasil dan participant tersimpan",
#                 "form_response_id": form_response.id
#             })

#         except Exception as e:
#             return JsonResponse({
#                 "status": "failed",
#                 "message": str(e)
#             })

#     return JsonResponse({
#         "status": "failed",
#         "message": "Invalid request"
#     })



@csrf_exempt
def upload_participant(request):

    if request.method != "POST":
        return JsonResponse({
            "status": "failed",
            "message": "Invalid request"
        })

    uploaded_file = request.FILES.get("excelFile")
    event_id = request.POST.get("event_id")

    if not uploaded_file:
        return JsonResponse({
            "status": "failed",
            "message": "File tidak ditemukan"
        })

    if not event_id:
        return JsonResponse({
            "status": "failed",
            "message": "Event tidak ditemukan"
        })

    try:
        event = Event.objects.get(id=event_id)
        status_default = StatusParticipant.objects.get(id=3)

        file_name = uploaded_file.name.lower()

        # =========================
        # READ FILE
        # =========================
        if file_name.endswith(".csv"):
            df = pd.read_csv(uploaded_file)

        elif file_name.endswith(".xlsx"):
            df = pd.read_excel(uploaded_file, engine="openpyxl")

        elif file_name.endswith(".xls"):
            df = pd.read_excel(uploaded_file, engine="xlrd")

        else:
            return JsonResponse({
                "status": "failed",
                "message": "Format file tidak didukung"
            })

        # =========================
        # CLEAN DATA
        # =========================
        def clean_value(val):
            bulan = [
                "Januari", "Februari", "Maret", "April",
                "Mei", "Juni", "Juli", "Agustus",
                "September", "Oktober", "November", "Desember"
            ]

            if isinstance(val, pd.Timestamp):
                dt = val.to_pydatetime()

                # 🔥 KONVERSI KE WITA (+1 jam dari WIB)
                dt = dt + timedelta(hours=1)

                return f"{dt.day} {bulan[dt.month-1]} {dt.year}, {dt.strftime('%H:%M')} WITA"

            elif isinstance(val, np.integer):
                return int(val)

            elif isinstance(val, np.floating):
                return float(val)

            elif pd.isna(val):
                return ""

            return val
        

        df = df.applymap(clean_value)
        excel_json = df.to_dict(orient="records")

        form_obj = Form.objects.first()

        saved_data = []

        # =========================
        # LOOP DATA
        # =========================
        df.columns = [
            re.sub(r"^\d+\.\s*", "",  # 🔥 hapus "01. ", "2. ", dll di depan
                re.sub(r"\s+", " ", str(col).replace(":", "").strip())
            )
            for col in df.columns
        ]

        name_col = find_column(df, [
            "name", "nama", "nama lengkap"
        ])

        email_col = find_column(df, [
            "email", "email address","alamat email"
        ])

        phone_col = find_column(df, [
            "phone", "no hp", "nohandphone",
            "nomor hp", "nomor telepon", "telepon"
        ])

        location_col = find_column(df, [
            "location", "alamat", "kota", "address"
        ])


        for _, row in df.iterrows():

            name = row[name_col] if name_col else ""
            email = row[email_col] if email_col else ""
            phone = row[phone_col] if phone_col else ""
            location = row[location_col] if location_col else ""

            name = str(name).strip()
            phone = str(phone).strip()

            # hapus .0 dari excel
            phone = re.sub(r"\.0$", "", phone)

            if not any([name, email, phone, location]):
                continue

            # =========================
            # CHECK EXISTING PARTICIPANT
            # =========================
            participant = Participant.objects.filter(
                name__iexact=name,
                phone=phone,
                id_event=event
            ).first()

            # =========================
            # UPDATE JIKA ADA
            # =========================
            if participant:

                participant.email = email
                participant.location = location
                participant.save()

            # =========================
            # CREATE BARU JIKA TIDAK ADA
            # =========================
            else:

                participant = Participant.objects.create(
                    name=name,
                    email=email,
                    phone=phone,
                    location=location,
                    id_event=event,
                    status_user_id=status_default
                )

            # 🔥 SIMPAN JSON + PARTICIPANT ID
            response_data = {
                "participant_id": participant.id,
                "participant_name": participant.name,
                "participant_email": participant.email,
                "participant_phone": participant.phone,
                 "excel_data": {
                    re.sub(r"^\d+\.\s*", "", key): value
                    for key, value in row.to_dict().items()
                }
            }

            form_response = FormResponse.objects.create(
                form=form_obj,
                data_response=response_data,
                participant=participant
            )

            saved_data.append({
                "participant_id": participant.id,
                "form_response_id": form_response.id
            })

        create_and_send_qremail(request,event_id)

        return JsonResponse({
            "status": "success",
            "message": "Upload berhasil",
            "total_uploaded": len(saved_data),
            "data": saved_data
        })

    except Exception as e:
        return JsonResponse({
            "status": "failed",
            "message": str(e)
        })
    


@csrf_exempt
def download_participant(request):

    if request.method != "GET":
        return JsonResponse({
            "status": "failed",
            "message": "Invalid request"
        })

    event_id = request.GET.get("event_id")

    if not event_id:
        return JsonResponse({
            "status": "failed",
            "message": "Event tidak ditemukan"
        })

    try:
        event = Event.objects.get(id=event_id)

        status_present = StatusParticipant.objects.get(
            status_name__iexact="Present"
        )

        participants = Participant.objects.filter(
            id_event=event,
            status_user_id=status_present
        ).values_list("name", "email", "phone")

        # ✅ DataFrame multi kolom
        df = pd.DataFrame(participants, columns=["Name", "Email", "Phone"])

        # ✅ Tambah kolom status
        df["Status"] = "Check In"

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        response['Content-Disposition'] = 'attachment; filename=list_name_participation_lucky_wheel.xlsx'

        df.to_excel(response, index=False, header=True)

        return response

    except StatusParticipant.DoesNotExist:
        return JsonResponse({
            "status": "failed",
            "message": "Status 'Present' tidak ditemukan"
        })

    except Event.DoesNotExist:
        return JsonResponse({
            "status": "failed",
            "message": "Event tidak ditemukan"
        })

    except Exception as e:
        return JsonResponse({
            "status": "failed",
            "message": str(e)
        })

@csrf_exempt
def update_doorprize(request):

    if request.method != "POST":
        return JsonResponse({
            "status": "failed",
            "message": "Invalid request"
        })

    try:

        # =========================================
        # GET PARAMETER
        # =========================================

        event_id = request.POST.get('event_id')

        prize_id = request.POST.get('prize_id')

        participant_id = request.POST.get('participant_id')

        # =========================================
        # NORMALIZE PARTICIPANT ID
        # =========================================

        if participant_id:

            participant_id = str(participant_id).replace('.', '')

            try:
                participant_id = int(participant_id)

            except (ValueError, TypeError):

                return JsonResponse({
                    'status': False,
                    'message': 'Invalid participant ID'
                })

        # =========================================
        # VALIDATION
        # =========================================

        if not event_id:

            return JsonResponse({
                'status': False,
                'message': 'Event required'
            })

        if not prize_id:

            return JsonResponse({
                'status': False,
                'message': 'Prize required'
            })

        if not participant_id:

            return JsonResponse({
                'status': False,
                'message': 'Participant required'
            })

        # =========================================
        # GET OBJECT
        # =========================================

        event = Event.objects.filter(
            id=event_id
        ).first()

        if not event:

            return JsonResponse({
                'status': False,
                'message': 'Event not found'
            })

        prize = Doorprize.objects.filter(
            id=prize_id,
            id_event=event,
            status='ACTIVE'
        ).first()

        if not prize:

            return JsonResponse({
                'status': False,
                'message': 'Prize already used or not found'
            })

        participant = Participant.objects.select_related(
            'status_user_id'
        ).filter(
            id=participant_id,
            id_event=event,
            status_user_id__status_name__iexact='Present'
        ).first()

        if not participant:

            return JsonResponse({
                'status': False,
                'message': 'Participant not eligible'
            })

        # =========================================
        # CHECK ALREADY WINNER
        # =========================================

        already_winner = DoorprizeWinner.objects.filter(
            id_event=event,
            id_participant=participant
        ).exists()

        if already_winner:

            return JsonResponse({
                'status': False,
                'message': 'Participant already won'
            })

        # =========================================
        # SAVE WINNER HISTORY
        # =========================================

        DoorprizeWinner.objects.create(

            id_event=event,

            id_doorprize=prize,

            id_participant=participant

        )

        participant.is_winner = True

        participant.save(
            update_fields=['is_winner']
        )

        # =========================================
        # UPDATE PRIZE
        # =========================================

        prize.winner = participant

        prize.status = 'DONE'

        prize.save()

        # =========================================
        # SUCCESS RESPONSE
        # =========================================

        return JsonResponse({

            'status': True,

            'message': 'Winner updated successfully',

            'winner_name': participant.name,

            'winner_id': participant.id,

            'prize_name': prize.prize_name,

            'event_name': event.name,

        })

    except Exception as e:

        return JsonResponse({

            'status': False,

            'message': str(e)

        })
    

@csrf_exempt
def add_prize(request):
    if request.method != "POST":
        return JsonResponse({
            "status": "failed",
            "message": "Invalid request"
        })

    try:

        event_id = request.POST.get("event_id")
        prize_name = request.POST.get("prize_name")
        qty = request.POST.get("qty", 1)

        if not prize_name:

            return JsonResponse({
                "status": "failed",
                "message": "Prize name required"
            })

        event = Event.objects.get(id=event_id)

        doorprize = Doorprize.objects.create(
            id_event=event,
            prize_name=prize_name,
            qty=qty,
            status="ACTIVE"
        )

        return JsonResponse({
            "status": "success",
            "id": doorprize.id,
            "prize_name": doorprize.prize_name
        })

    except Exception as e:

        return JsonResponse({
            "status": "failed",
            "message": str(e)
        })
    


@csrf_exempt
def add_participant_wheel(request):

    if request.method != "POST":

        return JsonResponse({
            "status": "failed",
            "message": "Invalid request"
        })

    try:

        participant_name = request.POST.get("participant_name")
        event_id = request.POST.get("event_id")

        # =========================================
        # VALIDASI
        # =========================================

        if not participant_name or not event_id:

            return JsonResponse({
                "status": "failed",
                "message": "Participant name dan Event wajib diisi"
            })

        # =========================================
        # GET EVENT
        # =========================================

        event = Event.objects.get(
            id=event_id
        )

        # =========================================
        # GET STATUS PRESENT
        # =========================================

        present_status = StatusParticipant.objects.get(
            status_name__iexact="Present"
        )

        # =========================================
        # CEK PARTICIPANT SUDAH ADA ATAU BELUM
        # =========================================

        participant = Participant.objects.filter(
            name__iexact=participant_name.strip(),
            id_event=event
        ).first()

        # =========================================
        # CREATE PARTICIPANT BARU
        # =========================================

        if not participant:

            participant = Participant.objects.create(
                id_event=event,
                name=participant_name.strip(),
                status_user_id=present_status,
                is_winner=False
            )

        # =========================================
        # VALIDASI STATUS PRESENT
        # =========================================

        if not participant.status_user_id:

            return JsonResponse({
                "status": "failed",
                "message": "Participant belum memiliki status"
            })

        if participant.status_user_id.status_name.lower() != "present":

            return JsonResponse({
                "status": "failed",
                "message": f"{participant.name} belum PRESENT"
            })

        # =========================================
        # VALIDASI SUDAH MENANG
        # CEK DARI is_winner
        # =========================================

        if participant.is_winner:

            return JsonResponse({
                "status": "failed",
                "message": f"{participant.name} sudah pernah menang"
            })

        # =========================================
        # VALIDASI DOUBLE DI WHEEL
        # =========================================

        # =========================================
        # SUCCESS
        # =========================================

        return JsonResponse({

            "status": "success",

            "participant": {
                "id": participant.id,
                "name": participant.name,
            }

        })

    except StatusParticipant.DoesNotExist:

        return JsonResponse({
            "status": "failed",
            "message": "Status PRESENT tidak ditemukan"
        })

    except Event.DoesNotExist:

        return JsonResponse({
            "status": "failed",
            "message": "Event tidak ditemukan"
        })

    except Exception as e:

        return JsonResponse({
            "status": "failed",
            "message": str(e)
        })
    
@csrf_exempt
def delete_participant_wheel(request):

    if request.method == "POST":

        participant_id = request.POST.get("participant_id")

        participant = Participant.objects.filter(
            id=participant_id
        ).first()

        if not participant:

            return JsonResponse({
                "status": False,
                "message": "Participant not found"
            })

        participant.is_winner = True
        participant.save()

        return JsonResponse({
            "status": True,
            "message": "Participant deleted"
        })

    return JsonResponse({
        "status": False,
        "message": "Invalid request"
    })

@csrf_exempt
def delete_prize(request):

    if request.method == "POST":

        prize_id = request.POST.get("prize_id")

        prize = Doorprize.objects.filter(
            id=prize_id
        ).first()

        if not prize:

            return JsonResponse({
                "status": False,
                "message": "Prize not found"
            })

        prize.delete()

        return JsonResponse({
            "status": True,
            "message": "Prize deleted"
        })

    return JsonResponse({
        "status": False,
        "message": "Invalid request"
    })




@csrf_exempt
def export_winner_excel(request, event_id):

    winners = DoorprizeWinner.objects.filter(
        id_event_id=event_id
    )

    workbook = openpyxl.Workbook()

    worksheet = workbook.active

    worksheet.title = "Winner History"

    headers = [
        "No",
        "Participant",
        "Prize"
    ]

    for col_num, header in enumerate(headers, 1):

        cell = worksheet.cell(
            row=1,
            column=col_num
        )

        cell.value = header

        cell.font = Font(bold=True)

    for index, winner in enumerate(winners, start=2):

        worksheet.cell(
            row=index,
            column=1
        ).value = index - 1

        worksheet.cell(
            row=index,
            column=2
        ).value = winner.id_participant.name

        worksheet.cell(
            row=index,
            column=3
        ).value = winner.id_doorprize.prize_name

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response[
        "Content-Disposition"
    ] = 'attachment; filename="winner_history.xlsx"'

    workbook.save(response)

    return response


@csrf_exempt
def search_participant(request):

    keyword = request.GET.get('keyword', '')
    event_id = request.GET.get('event_id')

    participants = Participant.objects.filter(
        id_event_id=event_id,
        status_user_id__status_name__iexact='Present',
        is_winner=False,
        name__icontains=keyword
    ).order_by('name')

    data = []

    for participant in participants:

        data.append({
            'id': participant.id,
            'name': participant.name,
        })

    return JsonResponse({
        'status': True,
        'participants': data
    })


@csrf_exempt
def search_prize(request):

    keyword = request.GET.get('keyword', '')
    event_id = request.GET.get('event_id')

    prizes = Doorprize.objects.filter(
        id_event_id=event_id,
        status='ACTIVE',
        prize_name__icontains=keyword
    ).order_by('id')

    data = []

    for prize in prizes:

        data.append({
            'id': prize.id,
            'prize_name': prize.prize_name,
        })

    return JsonResponse({
        'status': True,
        'prizes': data
    })